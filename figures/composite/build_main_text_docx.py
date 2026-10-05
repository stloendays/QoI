"""Main text as a submission-style Word document, with native (editable) Word equations.

Sources (read only): paper/MANUSCRIPT.md, paper/FIGURE_CAPTIONS.md, figures/composite/fig*/Fig*.png.
Output: figures/composite/release/QSQ_main_text_<date>.docx (release/ is gitignored: unpublished manuscript).

What it does beyond build_manuscript.py's quick export:
  * real Title / Heading 1-3 styles (navigation pane, automatic table of contents);
  * every $...$ and $$...$$ becomes Office Math (LaTeX -> MathML via latex2mathml -> OMML via MML2OMML.XSL);
  * figures and legends follow the references (submission layout), one per page, in order of first
    citation; panel letters lowercased as in the figures;
  * A4, Arial 11 pt, 1.5 line spacing, continuous line numbers, page numbers;
  * references as hanging-indent paragraphs with their original numbers.

    D:/Research/CatalystForge/.venv/Scripts/python.exe build_main_text_docx.py
"""
import copy
import os
import re
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_manuscript as B  # noqa: E402  (parse_md, captions, lower_panels, embed_png, OUT, PAPER, DATE)

XSL_PATH = r"D:\WPS Office\12.1.0.28802\office6\addons\kcoworkdockpanel\service\skills\wps-draft-by-template\scripts\oxir\MML2OMML.XSL"
# This copy ships with its root template commented out (it is called in mode="mml"); a two-line driver restores it.
_DRIVER = """<xsl:stylesheet version="1.0" xmlns:xsl="http://www.w3.org/1999/XSL/Transform"
  xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
  <xsl:import href="%s"/>
  <xsl:template match="/"><m:oMath><xsl:apply-templates mode="mml"/></m:oMath></xsl:template>
</xsl:stylesheet>""" % ("file:///" + XSL_PATH.replace("\\", "/").replace(" ", "%20"))
MML2OMML = etree.XSLT(etree.fromstring(_DRIVER.encode("utf-8")))
FONT, INK = "Arial", RGBColor(0x1B, 0x1B, 0x1B)
TEXT_W_CM = 17.0
TITLE = None


# ---- math ------------------------------------------------------------------------------------------------
def _clean(tex):
    tex = tex.strip()
    tex = tex.replace(r"\%", "%")
    tex = re.sub(r"\\text\{\\AA\}|\\mathrm\{\\AA\}|\\AA\b", r"\\mathrm{Å}", tex)          # angstrom
    tex = re.sub(r"\\(?:,|;|:|!)", lambda m: {r"\,": r"\,", r"\;": r"\;", r"\:": r"\;", r"\!": ""}[m.group(0)], tex)
    return tex


def omml(tex):
    """LaTeX -> <m:oMath> element (lxml), via MathML and Office's XSLT."""
    mml = etree.fromstring(latex_to_mathml(_clean(tex)).encode("utf-8"))
    om = MML2OMML(mml).getroot()
    if om.tag != qn("m:oMath"):
        found = om.find(".//" + qn("m:oMath"))
        om = found if found is not None else om
    _font_math_runs(om)
    return om


def _font_math_runs(om):
    """Cambria Math for every math run (WPS/Word otherwise fall back to the body font)."""
    for r in om.iter(qn("m:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = etree.SubElement(r, qn("w:rPr"))
            r.remove(rpr)
            r.insert(1 if r.find(qn("m:rPr")) is not None else 0, rpr)
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = etree.SubElement(rpr, qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), "Cambria Math")


# ---- inline text: **bold**, *italic*, `code`, $math$ ----------------------------------------------------
TOKEN = re.compile(r"(\$[^$]+\$|\*\*|(?<![\w*])\*(?=\S)|(?<=\S)\*(?![\w*])|`[^`]+`)")


def add_inline(par, s, size, bold=False, italic=False, color=INK):
    b, it = bold, italic
    for part in TOKEN.split(s):
        if not part:
            continue
        if part == "**":
            b = not b
        elif part == "*":
            it = not it
        elif part.startswith("$") and part.endswith("$") and len(part) > 2:
            par._p.append(omml(part[1:-1]))
        else:
            if part.startswith("`") and part.endswith("`"):
                part = part[1:-1]
            r = par.add_run(part)
            r.bold, r.italic = b, it
            r.font.size, r.font.name = Pt(size), FONT
            r.font.color.rgb = color
            r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


# ---- document setup ---------------------------------------------------------------------------------------
def style_font(st, size, bold=False, italic=False, color=INK, before=0, after=6, spacing=1.5, keep=False):
    st.font.name, st.font.size, st.font.bold, st.font.italic = FONT, Pt(size), bold, italic
    st.font.color.rgb = color
    rpr = st.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    pf.keep_with_next = keep


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        r = p.add_run()
        r.font.size, r.font.name = Pt(9), FONT
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)


def line_numbers(section):
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "283")
    section._sectPr.append(ln)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    S = doc.styles
    style_font(S["Normal"], 11, after=6)
    style_font(S["Title"], 17, bold=True, after=14, spacing=1.15)
    style_font(S["Heading 1"], 14, bold=True, before=18, after=6, spacing=1.15, keep=True)
    style_font(S["Heading 2"], 12, bold=True, before=14, after=4, spacing=1.15, keep=True)
    style_font(S["Heading 3"], 11, bold=True, italic=False, before=10, after=3, spacing=1.15, keep=True)
    style_font(S["Caption"], 9, after=12, spacing=1.2)
    # the Title style has a bottom border in the default template; drop it
    pbdr = S["Title"].element.pPr.find(qn("w:pBdr")) if S["Title"].element.pPr is not None else None
    if pbdr is not None:
        pbdr.getparent().remove(pbdr)
    try:
        lnst = S["Line Number"]
    except KeyError:
        lnst = S.add_style("Line Number", WD_STYLE_TYPE.CHARACTER)
    lnst.font.name, lnst.font.size, lnst.font.color.rgb = FONT, Pt(7.5), RGBColor(0x55, 0x5A, 0x62)
    page_number_footer(sec)
    line_numbers(sec)
    doc.core_properties.title = TITLE or ""
    doc.core_properties.author = ""
    return doc


# ---- assembly ---------------------------------------------------------------------------------------------
def build():
    global TITLE
    md = open(os.path.join(B.PAPER, "MANUSCRIPT.md"), encoding="utf-8").read()
    caps = B.captions(os.path.join(B.PAPER, "FIGURE_CAPTIONS.md"), "Figure")
    blocks = B.parse_md(md)
    TITLE = next(b[1] for b in blocks if b[0] == "h1")
    doc = new_document()
    placed, in_refs, n_math = [], False, 0
    for b in blocks:
        k = b[0]
        if k == "h1":
            doc.add_paragraph(b[1], style="Title")
        elif k == "h2":
            in_refs = b[1].strip().lower() == "references"
            h = doc.add_paragraph(style="Heading 1"); add_inline(h, b[1], 14, bold=True)
        elif k == "h3":
            h = doc.add_paragraph(style="Heading 2"); add_inline(h, b[1], 12, bold=True)
        elif k == "eq":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para = OxmlElement("m:oMathPara")
            para.append(omml(b[1])); p._p.append(para); n_math += 1
        elif k == "p":
            text = B.lower_panels(b[1])
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, text, 11)
            n_math += len(re.findall(r"\$[^$]+\$", text))
            for n in re.findall(r"Fig(?:\.|ure)\s*(\d)", text):
                if n not in placed:
                    placed.append(n)
        elif k in ("ul", "ol"):
            for j, x in enumerate(b[1], start=1):
                if k == "ol" and in_refs:
                    p = doc.add_paragraph()
                    pf = p.paragraph_format
                    pf.left_indent, pf.first_line_indent = Cm(0.8), Cm(-0.8)
                    pf.space_after, pf.line_spacing = Pt(3), 1.15
                    add_inline(p, "%d.\t" % j, 9); add_inline(p, x, 9)
                    pf.tab_stops.add_tab_stop(Cm(0.8))
                else:
                    p = doc.add_paragraph(style="List Bullet" if k == "ul" else "List Number")
                    add_inline(p, x, 11)
                n_math += len(re.findall(r"\$[^$]+\$", x))
        elif k == "table":
            B.T.to_docx_table(doc, b) if hasattr(B.T, "to_docx_table") else None
    assert placed == [str(i) for i in range(1, 10)], placed         # cited in order 1..9
    # ---- figures and legends, one per page, after the references --------------------------------
    pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Figures", style="Heading 1")
    for j, n in enumerate(placed):
        png = B.embed_png(B.main_figure_path(n), "Fig%s" % n)
        fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.keep_with_next = True
        fp.paragraph_format.space_before = Pt(6 if j == 0 else 0)
        fp.paragraph_format.line_spacing = 1.0
        fp.add_run().add_picture(png, width=Cm(TEXT_W_CM))
        cp = doc.add_paragraph(style="Caption"); cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(cp, caps[n], 9)
        n_math += len(re.findall(r"\$[^$]+\$", caps[n]))
        if j < len(placed) - 1:
            cp.add_run().add_break(WD_BREAK.PAGE)
    path = os.path.join(B.OUT, "QSQ_main_text_%s.docx" % B.DATE)
    doc.save(path)
    return path, n_math


if __name__ == "__main__":
    os.makedirs(B.OUT, exist_ok=True)
    path, n = build()
    print("wrote %s  (%d math objects, %.1f MB)" % (path, n, os.path.getsize(path) / 1e6))

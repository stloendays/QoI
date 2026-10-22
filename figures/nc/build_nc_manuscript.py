"""Nature Communications manuscript as PDF + DOCX with the eight main figures embedded.

Reuses the parser of figures/composite/build_manuscript.py and changes:
  - equations: MathML in the PDF and native Word equations (OMML, omml.py) in the DOCX;
  - the figure mapping (8 main figures: composite Figs. 3 and 7 and figures/nc/fig1, fig4-fig8), and
  - the caption source (the '## Figure legends' section of paper/MANUSCRIPT.md, which is removed from the body because
    each figure is inserted with its legend after the paragraph that first cites it).

    D:/Research/CatalystForge/.venv/Scripts/python.exe build_nc_manuscript.py
Outputs: figures/nc/release/NC_manuscript_<date>.{html,pdf,docx} and MANIFEST_<date>.txt
"""
import hashlib
import html
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from latex2mathml.converter import convert as to_mathml

HERE = os.path.dirname(os.path.abspath(__file__))
COMP = os.path.join(os.path.dirname(HERE), "composite")
sys.path.insert(0, COMP)
import build_manuscript as B  # noqa: E402
sys.path.insert(0, HERE)
import omml as OM  # noqa: E402

REPO = B.REPO
B.OUT = os.path.join(HERE, "release")
B.EMBED = os.path.join(B.OUT, "_embed")

FIG = {
    "1": os.path.join(HERE, "fig1", "Fig1.png"),
    "2": os.path.join(COMP, "fig3", "Fig3.png"),
    "3": os.path.join(COMP, "fig7", "Fig7.png"),
    "4": os.path.join(HERE, "fig4", "Fig4.png"),
    "5": os.path.join(HERE, "fig5", "Fig5.png"),
    "6": os.path.join(HERE, "fig6", "Fig6.png"),
    "7": os.path.join(HERE, "fig7", "Fig7.png"),
    "8": os.path.join(HERE, "fig8", "Fig8.png"),
}


def split_legends(md):
    body, leg = md.split("## Figure legends", 1)
    caps = {}
    for para in [p.strip() for p in leg.strip().split("\n\n") if p.strip()]:
        m = re.match(r"\*\*Fig\. (\d+) \|", para)
        assert m, para[:60]
        caps[m.group(1)] = para
    return body, caps


SUP = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")


def superscripts(text):
    """Nature-style citation markers <sup>1–4,7</sup> -> Unicode superscripts (works in both the PDF and the DOCX)."""
    return re.sub(r"<sup>([^<]*)</sup>", lambda m: m.group(1).translate(SUP), text)


# ---- equations: MathML in the PDF (Chrome renders it natively), native OMML equations in the DOCX ---------------
MATH = re.compile(r"(?<![$\\])\$(?!\$)(.+?)(?<!\\)\$")


def prep(latex):
    """The few constructs latex2mathml does not read: \\AA inside \\text, escaped underscores in \\text."""
    return latex.replace(r"\text{\AA}", r"\text{Å}").replace(r"\AA", "Å").replace(r"\_", "_").strip()


def mathml(latex, display="inline"):
    """MathML with print conventions Chrome does not apply by itself: upright capital Greek, and no operator
    spacing around | and / (they delimit and divide here, they do not separate terms)."""
    m = to_mathml(prep(latex), display=display)
    m = re.sub(r"<mi>(&#x0039[1-9A-F];|&#x003A[0-9];)</mi>", r'<mi mathvariant="normal">\1</mi>', m)
    m = re.sub(r"<mo>(&#x0007C;|&#x0002F;)</mo>", r'<mo lspace="0" rspace="0">\1</mo>', m)
    m = re.sub(r'<mo stretchy="false">(&#x0007C;)</mo>', r'<mo stretchy="false" lspace="0" rspace="0">\1</mo>', m)
    return m


def styled(s):
    """Markdown run -> [(text, bold, italic, is_math)]; math spans are kept as LaTeX."""
    s = re.sub(r"`([^`]+)`", r"\1", s)
    pieces, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*", s):
        if m.start() > pos:
            pieces.append((s[pos:m.start()], False, False))
        pieces.append((m.group(1), True, False) if m.group(1) is not None else (m.group(2), False, True))
        pos = m.end()
    if pos < len(s):
        pieces.append((s[pos:], False, False))
    out = []
    for text, b, it in pieces:
        p = 0
        for m in MATH.finditer(text):
            if m.start() > p:
                out.append((text[p:m.start()], b, it, False))
            out.append((m.group(1), b, it, True))
            p = m.end()
        if p < len(text):
            out.append((text[p:], b, it, False))
    return out


def html_runs(s):
    o = []
    for t, b, it, is_math in styled(s):
        t = mathml(t) if is_math else html.escape(t)
        o.append(("<b>%s</b>" if b else "<i>%s</i>" if it else "%s") % t)
    return "".join(o)


def add_runs(par, s, size, bold_all=False):
    for t, b, it, is_math in styled(s):
        if is_math:
            OM.add_math(par, prep(t), size)
            continue
        r = par.add_run(t); r.font.size = Pt(size); r.font.name = "Arial"; r.bold = b or bold_all; r.italic = it


MATH_CSS = "math { font-family: 'Cambria Math', 'STIX Two Math', serif; font-size: 1.05em; } p.eq math { font-size: 1.12em; }\n"


def blocks_html(blocks):
    out = []
    for b in blocks:
        k = b[0]
        if k in ("h1", "h2", "h3"):
            out.append("<%s>%s</%s>" % (k, html_runs(b[1]), k))
        elif k == "eq":
            out.append("<p class='eq'>%s</p>" % mathml(b[1], display="block"))
        elif k == "p":
            out.append("<p>%s</p>" % html_runs(b[1]))
        elif k == "ol":
            out.append("<ol class='refs'>" + "".join("<li>%s</li>" % html_runs(x) for x in b[1]) + "</ol>")
        elif k == "ul":
            out.append("<ul>" + "".join("<li>%s</li>" % html_runs(x) for x in b[1]) + "</ul>")
        elif k == "fig":
            out.append("<figure><img src='%s'><figcaption>%s</figcaption></figure>" % (b[1].replace("\\", "/"), html_runs(b[2])))
        elif k == "table":
            _, header, align, rows = b
            out.append("<table><thead><tr>" + "".join("<th class='%s'>%s</th>" % (a, html_runs(h)) for h, a in zip(header, align))
                       + "</tr></thead><tbody>")
            for r in rows:
                r = r + [""] * (len(header) - len(r))
                out.append("<tr>" + "".join("<td class='%s'>%s</td>" % (a, html_runs(c)) for c, a in zip(r, align)) + "</tr>")
            out.append("</tbody></table>")
    return "\n".join(out)


def write_html_pdf(blocks, stem, title):
    h = "<!doctype html><meta charset='utf-8'><title>%s</title><style>%s%s</style><body>%s</body>" % (
        html.escape(title), B.CSS, MATH_CSS, blocks_html(blocks))
    hp = os.path.join(B.OUT, stem + ".html")
    open(hp, "w", encoding="utf-8").write(h)
    pdf = os.path.join(B.OUT, stem + ".pdf")
    if os.path.exists(pdf):
        os.remove(pdf)
    subprocess.run([B.CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--print-to-pdf=%s" % pdf,
                    "file:///" + hp.replace("\\", "/")], check=False, capture_output=True, timeout=600)
    return pdf


def write_docx(blocks, stem):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6); sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for b in blocks:
        k = b[0]
        if k in ("h1", "h2", "h3"):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt({"h1": 0, "h2": 16, "h3": 10}[k])
            p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
            add_runs(p, b[1], {"h1": 15, "h2": 11.5, "h3": 10}[k], bold_all=True)
        elif k == "eq":
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(6)
            OM.add_display_math(p, prep(b[1]), 10.5, text_width_cm=B.TEXT_W_CM)
        elif k == "p":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, b[1], 10)
        elif k in ("ul", "ol"):
            for x in b[1]:
                p = doc.add_paragraph(style="List Bullet" if k == "ul" else "List Number")
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, x, 10 if k == "ul" else 8.5)
        elif k == "fig":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.add_run().add_picture(b[1], width=Cm(B.TEXT_W_CM))
            c = doc.add_paragraph(); c.paragraph_format.space_after = Pt(12); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(c, b[2], 8.3)
        elif k == "table":
            B.T.to_docx_table(doc, b)
    path = os.path.join(B.OUT, stem + ".docx")
    doc.save(path)
    return path


def main_manuscript():
    md = superscripts(open(os.path.join(REPO, "paper", "MANUSCRIPT.md"), encoding="utf-8").read())
    body, caps = split_legends(md)
    assert set(caps) == set(FIG), (sorted(caps), sorted(FIG))
    blocks = [(b[0], B.lower_panels(b[1])) if b[0] == "p" else b for b in B.parse_md(body)]
    placed, out = set(), []
    for b in blocks:
        out.append(b)
        if b[0] == "p":
            for n in re.findall(r"Fig(?:\.|ure)s?\.? (\d)", b[1]):
                if n in FIG and n not in placed:
                    placed.add(n)
                    png = B.embed_png(FIG[n], "NCFig%s" % n)
                    out.append(("fig", png, caps[n]))
    assert placed == set(FIG), placed
    return out


def _docx_table_patch():
    """Same table writer that build_manuscript.main() attaches to the SI table module."""
    T = B.T
    if hasattr(T, "to_docx_table"):
        return
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    def to_docx_table(doc, b):
        _, header, align, rows = b
        t = doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, (h, a) in enumerate(zip(header, align)):
            c = t.cell(0, j); c.text = ""; p = c.paragraphs[0]; add_runs(p, h, 7.6, bold_all=True)
            if a == "r":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            T.cell_border(c, top=6, bottom=4)
        for i, r in enumerate(rows, start=1):
            r = r + [""] * (len(header) - len(r))
            for j, (v, a) in enumerate(zip(r, align)):
                c = t.cell(i, j); c.text = ""; p = c.paragraphs[0]; add_runs(p, v, 7.6)
                p.paragraph_format.space_after = Pt(0)
                if a == "r":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                T.cell_border(c, bottom=6 if i == len(rows) else None)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    T.to_docx_table = to_docx_table


def main():
    os.makedirs(B.OUT, exist_ok=True)
    _docx_table_patch()
    title = open(os.path.join(REPO, "paper", "MANUSCRIPT.md"), encoding="utf-8").readline().lstrip("# ").strip()
    stem = "NC_manuscript_%s" % B.DATE
    blocks = main_manuscript()
    pdf = write_html_pdf(blocks, stem, title)
    docx = write_docx(blocks, stem)
    man = ["%s  %s  %d bytes" % (hashlib.sha256(open(p, "rb").read()).hexdigest(), os.path.basename(p), os.path.getsize(p))
           for p in (pdf, docx)]
    open(os.path.join(B.OUT, "MANIFEST_%s.txt" % B.DATE), "w", encoding="utf-8").write("\n".join(man) + "\n")
    print("\n".join(man))


if __name__ == "__main__":
    main()
